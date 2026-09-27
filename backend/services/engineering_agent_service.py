"""Engineering drawing & documentation agent: review, Q&A, comparison, templates and document generation."""

import io
import json
import os
import re
import shutil
import uuid
from typing import Any, Dict, List, Optional

from langgraph.graph import END, StateGraph
from openai import OpenAI
from typing_extensions import TypedDict

from core.config import settings
from utils.logger import logger
from .engineering_loader import (
    SUPPORTED_EXTENSIONS, EngineeringLoader, crop_for_vision, page_images_for_vision, text_layer_summary,
)
from .engineering_prompts import (
    ASK_SYSTEM, COMPARE_SYSTEM, CROSS_CHECK_SYSTEM, DATUM_VERIFY_SYSTEM, EXTRACT_SYSTEM, GENERATE_SYSTEM, REVIEW_SYSTEM,
    REVISION_AUDIT_SYSTEM, STANDARDS,
    TEMPLATE_IMPORT_SYSTEM, TEMPLATE_REVIEW_SYSTEM,
)
from .engineering_rules import diff_extractions, run_rule_checks
from .prompt_guard import UNTRUSTED_NOTICE, detect_injection, injection_finding
from .engineering_store import EngineeringStore
from .file_safety import validate_upload
from .usage_tracker import record_usage


FIELD_TYPES = {"text", "longtext", "date", "table"}
SEVERITY_ORDER = {"critical": 0, "major": 1, "minor": 2, "info": 3}
TILE_LABELS = ["full sheet", "top-left tile", "top-right tile", "bottom-left tile", "bottom-right tile"]


class ReviewState(TypedDict):
    """State for the drawing review graph."""
    drawing_id: str
    filename: str
    file_path: str
    standard: str
    loaded: Dict[str, Any]
    extraction: Dict[str, Any]
    rule_checks: List[Dict[str, Any]]
    consistency_issues: List[Dict[str, Any]]
    review: Dict[str, Any]


class EngineeringAgentService:
    """Orchestrates the engineering agents; the review workflow is a LangGraph graph."""

    def __init__(self):
        """Initialize the client, store and review graph."""
        # Vision calls with several high-detail images can take a while; still retry stalled connections
        self.client = OpenAI(api_key=settings.OPENAI_API_KEY, base_url=settings.OPENAI_BASE_URL or None, max_retries=4, timeout=180)
        self.model = settings.ENGINEERING_MODEL
        self.store = EngineeringStore()
        self.loader = EngineeringLoader()
        self.files_dir = os.path.join(settings.ENGINEERING_DATA_DIR, "files")
        os.makedirs(self.files_dir, exist_ok=True)
        self.review_graph = self._build_review_graph()

    # ------------------------------------------------------------------ LLM helpers

    def _chat_json(
        self, system: str, content: Any, max_tokens: int = 4000, temperature: float = 0.1
    ) -> Dict[str, Any]:
        """Call the model in JSON mode; content is a string or a list of multimodal parts."""
        response = self.client.chat.completions.create(
            model=self.model,
            messages=[{"role": "system", "content": system + UNTRUSTED_NOTICE}, {"role": "user", "content": content}],
            response_format={"type": "json_object"},
            temperature=temperature,
            max_tokens=max_tokens,
        )
        record_usage(self.model, response.usage, "engineering")
        return json.loads(response.choices[0].message.content or "{}")

    @staticmethod
    def _sheet_images(pages: List[Dict[str, Any]], label_prefix: str = "") -> List[Dict[str, Any]]:
        """Multimodal parts for each sheet: full sheet plus four zoomed tiles, each labelled."""
        parts: List[Dict[str, Any]] = []
        for page in pages:
            for label, image_b64 in zip(TILE_LABELS, page_images_for_vision(page["image_path"])):
                parts.append({"type": "text", "text": f"{label_prefix}Sheet {page['number']} – {label}:"})
                parts.append({
                    "type": "image_url",
                    "image_url": {"url": f"data:image/png;base64,{image_b64}", "detail": "high"},
                })
        return parts

    # ------------------------------------------------------------------ drawing review

    def review_upload(self, filename: str, content: bytes, standard: str, owner_id: Optional[str] = None) -> Dict[str, Any]:
        """Save an uploaded drawing and run the full review workflow on it."""
        extension = os.path.splitext(filename)[1].lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type '{extension}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
        if len(content) > settings.ENGINEERING_MAX_FILE_MB * 1024 * 1024:
            raise ValueError(f"File is larger than {settings.ENGINEERING_MAX_FILE_MB} MB")
        validate_upload(filename, content, SUPPORTED_EXTENSIONS)  # raises UnsafeFileError (a ValueError)

        drawing_id = str(uuid.uuid4())
        directory = os.path.join(self.files_dir, drawing_id)
        os.makedirs(directory, exist_ok=True)
        # Stored under a generated name so user-supplied filenames never touch the filesystem path
        file_path = os.path.join(directory, f"original{extension}")
        with open(file_path, "wb") as f:
            f.write(content)
        self.store.create_drawing(drawing_id, os.path.basename(filename), file_path, standard, owner_id)
        return self.run_review(drawing_id, standard)

    def run_review(self, drawing_id: str, standard: str) -> Dict[str, Any]:
        """Run (or re-run) the review graph for a stored drawing."""
        drawing = self.store.get_drawing(drawing_id)
        if drawing is None:
            raise KeyError(drawing_id)
        self.store.update_drawing(drawing_id, status="processing", error=None, standard=standard)
        try:
            self.review_graph.invoke({
                "drawing_id": drawing_id,
                "filename": drawing["filename"],
                "file_path": drawing["file_path"],
                "standard": standard,
                "loaded": {},
                "extraction": {},
                "rule_checks": [],
                "consistency_issues": [],
                "review": {},
            })
            self.store.update_drawing(drawing_id, status="completed")
        except Exception as e:
            logger.error(f"Engineering review failed for {drawing['filename']}: {str(e)}")
            self.store.update_drawing(drawing_id, status="failed", error=str(e)[:500])
        return self.get_drawing(drawing_id)

    def _build_review_graph(self):
        """load -> extract -> verify_datums -> rule_check -> cross_check -> review -> finalize."""

        def load(state: ReviewState) -> ReviewState:
            loaded = self.loader.load(state["file_path"], state["filename"], os.path.dirname(state["file_path"]))
            state["loaded"] = loaded
            self.store.update_drawing(
                state["drawing_id"],
                source_type=loaded["source_type"],
                page_count=loaded["page_count"],
                pages=loaded["pages"],
                cad=loaded["cad"],
            )
            return state

        def extract(state: ReviewState) -> ReviewState:
            loaded = state["loaded"]
            content: List[Dict[str, Any]] = [{
                "type": "text",
                "text": f"File: {state['filename']} ({loaded['source_type']}, {loaded['page_count']} sheet(s))\n\n"
                        f"Text layer with positions (x from left, y from top):\n{text_layer_summary(loaded['pages'])}"
                        + (f"\n\nExact CAD data (units, block attributes, dimensions):\n"
                           f"{json.dumps(_cad_for_prompt(loaded['cad']))[:8000]}" if loaded["cad"] else ""),
            }]
            content += self._sheet_images(loaded["pages"])
            state["extraction"] = self._chat_json(EXTRACT_SYSTEM, content, max_tokens=6000)
            self.store.update_drawing(state["drawing_id"], extraction=state["extraction"])
            return state

        def verify_datums(state: ReviewState) -> ReviewState:
            # Vision models sometimes "see" a datum the design would plausibly have; confirm each claimed
            # datum symbol on a zoomed crop and drop the ones that cannot be found
            pages = {page["number"]: page for page in state["loaded"]["pages"]}
            verified, rejected = [], []
            for datum in state["extraction"].get("datum_features") or []:
                if not isinstance(datum, dict) or not datum.get("letter"):
                    continue
                page = pages.get(int(datum.get("page") or 1), state["loaded"]["pages"][0])
                try:
                    crop = crop_for_vision(page["image_path"], float(datum.get("x") or 50), float(datum.get("y") or 50))
                    result = self._chat_json(DATUM_VERIFY_SYSTEM, [
                        {"type": "text", "text": f"Claim: a datum feature symbol for datum '{datum['letter']}' is "
                                                 f"here, attached to {datum.get('attached_to')}. Is it present?"},
                        {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{crop}", "detail": "high"}},
                    ], max_tokens=300)
                except Exception as e:
                    logger.warning(f"Datum verification failed for {datum['letter']}: {str(e)}")
                    result = {"present": True, "reason": "verification unavailable"}
                (verified if result.get("present") else rejected).append({**datum, "verification": result.get("reason")})
            state["extraction"]["datum_features"] = verified
            state["extraction"]["unverified_datum_claims"] = rejected
            self.store.update_drawing(state["drawing_id"], extraction=state["extraction"])
            return state

        def rule_check(state: ReviewState) -> ReviewState:
            state["rule_checks"] = run_rule_checks(
                state["extraction"], state["loaded"]["page_count"], state["loaded"]["cad"]
            )
            self.store.update_drawing(state["drawing_id"], rule_checks=state["rule_checks"])
            return state

        def cross_check(state: ReviewState) -> ReviewState:
            # Specialist pass for contradictions between views and broken references, which a single
            # broad review tends to miss
            content: List[Dict[str, Any]] = [{
                "type": "text",
                "text": f"Extracted drawing data:\n{json.dumps(state['extraction'])[:15000]}\n\n"
                        f"Text layer:\n{text_layer_summary(state['loaded']['pages'], 8000)}",
            }]
            content += self._sheet_images(state["loaded"]["pages"])
            result = self._chat_json(CROSS_CHECK_SYSTEM, content, max_tokens=2500)
            state["consistency_issues"] = [i for i in result.get("issues") or [] if isinstance(i, dict)]
            return state

        def review(state: ReviewState) -> ReviewState:
            standard = STANDARDS[state["standard"]]
            failed_rules = [r for r in state["rule_checks"] if r["status"] == "fail"]
            content: List[Dict[str, Any]] = [{
                "type": "text",
                "text": f"Standard: {standard['label']}\nChecklist:\n"
                        + "\n".join(f"- {item}" for item in standard["checklist"])
                        + f"\n\nExtracted drawing data:\n{json.dumps(state['extraction'])[:15000]}"
                        + f"\n\nAutomated rule-check failures:\n{json.dumps(failed_rules) if failed_rules else 'none'}"
                        + f"\n\nConsistency checker issues:\n"
                          f"{json.dumps(state['consistency_issues']) if state['consistency_issues'] else 'none'}"
                        + f"\n\nText layer:\n{text_layer_summary(state['loaded']['pages'], 8000)}",
            }]
            content += self._sheet_images(state["loaded"]["pages"])
            state["review"] = self._chat_json(REVIEW_SYSTEM, content, max_tokens=6000)
            return state

        def finalize(state: ReviewState) -> ReviewState:
            result = state["review"]
            findings = [f for f in result.get("findings") or [] if isinstance(f, dict)]
            drawing_text = " ".join(
                block.get("text", "") for page in state["loaded"].get("pages") or [] for block in page.get("text") or []
                if isinstance(block, dict)
            )
            excerpts = detect_injection(drawing_text)
            if excerpts:
                findings.append(injection_finding(excerpts))
                result["injection_warnings"] = excerpts
            for finding in findings:
                if finding.get("severity") not in SEVERITY_ORDER:
                    finding["severity"] = "minor"
            findings.sort(key=lambda f: SEVERITY_ORDER[f["severity"]])
            result["findings"] = findings
            # Guardrail: the verdict can never be more lenient than the worst finding
            severities = {f["severity"] for f in findings}
            if "critical" in severities:
                result["verdict"] = "rejected"
            elif "major" in severities and result.get("verdict") == "approved":
                result["verdict"] = "approved_with_comments"
            elif result.get("verdict") not in ("approved", "approved_with_comments", "rejected"):
                result["verdict"] = "approved_with_comments"
            result["standard"] = state["standard"]
            result["consistency_issues"] = state["consistency_issues"]
            self.store.update_drawing(state["drawing_id"], review=result)
            return state

        workflow = StateGraph(ReviewState)
        workflow.add_node("load", load)
        workflow.add_node("extract", extract)
        workflow.add_node("verify_datums", verify_datums)
        workflow.add_node("rule_check", rule_check)
        workflow.add_node("cross_check", cross_check)
        workflow.add_node("review_drawing", review)
        workflow.add_node("finalize", finalize)
        workflow.set_entry_point("load")
        workflow.add_edge("load", "extract")
        workflow.add_edge("extract", "verify_datums")
        workflow.add_edge("verify_datums", "rule_check")
        workflow.add_edge("rule_check", "cross_check")
        workflow.add_edge("cross_check", "review_drawing")
        workflow.add_edge("review_drawing", "finalize")
        workflow.add_edge("finalize", END)
        return workflow.compile()

    def get_drawing(self, drawing_id: str) -> Optional[Dict[str, Any]]:
        """Return a drawing record for the API (without filesystem paths)."""
        drawing = self.store.get_drawing(drawing_id)
        if drawing is None:
            return None
        drawing.pop("file_path", None)
        drawing["pages"] = [
            {k: v for k, v in page.items() if k != "image_path"} for page in drawing.get("pages") or []
        ]
        if drawing.get("cad"):
            drawing["cad"] = _cad_for_prompt(drawing["cad"])
        return drawing

    def page_image_path(self, drawing_id: str, page_number: int) -> Optional[str]:
        drawing = self.store.get_drawing(drawing_id)
        for page in (drawing or {}).get("pages") or []:
            if page["number"] == page_number:
                return page["image_path"]
        return None

    def delete_drawing(self, drawing_id: str) -> bool:
        if self.store.get_drawing(drawing_id) is None:
            return False
        shutil.rmtree(os.path.join(self.files_dir, drawing_id), ignore_errors=True)
        self.store.delete_drawing(drawing_id)
        return True

    # ------------------------------------------------------------------ Q&A

    def ask(self, drawing_id: str, question: str, history: List[Dict[str, str]]) -> Dict[str, Any]:
        """Answer a question about a reviewed drawing."""
        drawing = self.store.get_drawing(drawing_id)
        if drawing is None or not drawing.get("pages"):
            raise KeyError(drawing_id)
        review = drawing.get("review") or {}
        conversation = "\n".join(f"{turn['role']}: {turn['content']}" for turn in history[-6:])
        content: List[Dict[str, Any]] = [{
            "type": "text",
            "text": f"Extracted drawing data:\n{json.dumps(drawing.get('extraction') or {})[:12000]}\n\n"
                    f"Review findings:\n{json.dumps(review.get('findings') or [])[:6000]}\n\n"
                    f"Text layer:\n{text_layer_summary(drawing['pages'], 6000)}\n\n"
                    + (f"Conversation so far:\n{conversation}\n\n" if conversation else "")
                    + f"Question: {question}",
        }]
        content += self._sheet_images(drawing["pages"][:2])
        return {"answer": self._chat_json(ASK_SYSTEM, content, max_tokens=1500).get("answer", "")}

    # ------------------------------------------------------------------ comparison

    def compare(self, a_id: str, b_id: str, owner_id: Optional[str] = None) -> Dict[str, Any]:
        """Compare two reviewed drawings/documents and store the result."""
        a, b = self.store.get_drawing(a_id), self.store.get_drawing(b_id)
        if a is None or b is None:
            raise KeyError("drawing not found")
        if not a.get("extraction") or not b.get("extraction"):
            raise ValueError("Both drawings must be reviewed before they can be compared")

        differences = diff_extractions(a["extraction"], b["extraction"])
        content: List[Dict[str, Any]] = [{
            "type": "text",
            "text": f"A = {a['filename']}\nB = {b['filename']}\n\n"
                    f"Extracted data A:\n{json.dumps(_comparable(a['extraction']))[:12000]}\n\n"
                    f"Extracted data B:\n{json.dumps(_comparable(b['extraction']))[:12000]}\n\n"
                    f"Automated difference list:\n{json.dumps(differences)[:8000]}",
        }]
        content += self._sheet_images(a["pages"][:1], "A – ")
        content += self._sheet_images(b["pages"][:1], "B – ")
        result = self._chat_json(COMPARE_SYSTEM, content, max_tokens=5000)
        # The model sometimes returns structured values (e.g. a revision-table row); the UI expects text
        result["changes"] = [
            {**change, "a": _as_text(change.get("a")), "b": _as_text(change.get("b")), "item": _as_text(change.get("item"))}
            for change in result.get("changes") or [] if isinstance(change, dict)
        ]
        result["inconsistencies"] = [i for i in result.get("inconsistencies") or [] if isinstance(i, dict)]
        result["revision_control"] = self._audit_revision_control(a["extraction"], b["extraction"], result)
        result["automated_differences"] = differences
        result["a"] = {"id": a_id, "filename": a["filename"]}
        result["b"] = {"id": b_id, "filename": b["filename"]}
        result["id"] = self.store.create_comparison(a_id, b_id, result, owner_id)
        return result

    def _audit_revision_control(
        self, a: Dict[str, Any], b: Dict[str, Any], comparison: Dict[str, Any]
    ) -> Dict[str, Any]:
        """Check each substantive change against the revision-table entries added in B."""
        revision_a = ((a.get("title_block") or {}).get("revision") or "").strip().upper()
        revision_b = ((b.get("title_block") or {}).get("revision") or "").strip().upper()
        incremented = None
        if revision_a and revision_b:
            incremented = (len(revision_b), revision_b) > (len(revision_a), revision_a)

        rows_a = {json.dumps(row, sort_keys=True) for row in a.get("revision_table") or []}
        new_entries = [row for row in b.get("revision_table") or [] if json.dumps(row, sort_keys=True) not in rows_a]
        substantive = [
            change for change in comparison.get("changes") or []
            if isinstance(change, dict) and change.get("impact") != "documentation_only"
        ]
        audit = []
        if substantive:
            audit = self._chat_json(
                REVISION_AUDIT_SYSTEM,
                f"New revision-table entries in B:\n{json.dumps(new_entries)}\n\n"
                f"Changes:\n{json.dumps([{'item': c.get('item'), 'a': c.get('a'), 'b': c.get('b')} for c in substantive])}",
                max_tokens=1500,
            ).get("audit") or []
        unrecorded = [
            f"{entry.get('item')}" for entry in audit if isinstance(entry, dict) and not entry.get("recorded")
        ]
        if not new_entries and substantive:
            unrecorded = [str(c.get("item")) for c in substantive]
        if incremented is False and substantive:
            comment = f"Substantive changes were made but the revision was not incremented ({revision_a} → {revision_b})."
        elif unrecorded:
            comment = (f"{len(unrecorded)} of {len(substantive)} substantive change(s) are not described in the "
                       f"revision table entries added in revision {revision_b or 'B'}.")
        else:
            comment = "Every substantive change is described in the new revision-table entries."
        return {
            "revision_incremented": incremented,
            "changes_recorded_in_revision_table": not unrecorded if substantive else None,
            "unrecorded_changes": unrecorded,
            "new_revision_entries": new_entries,
            "audit": audit,
            "comment": comment,
        }

    # ------------------------------------------------------------------ templates

    def review_template(self, template_id: str) -> Dict[str, Any]:
        """Have the agent critique a template and propose an improved version."""
        template = self.store.get_template(template_id)
        if template is None:
            raise KeyError(template_id)
        result = self._chat_json(TEMPLATE_REVIEW_SYSTEM, json.dumps(_template_payload(template)), max_tokens=5000)
        if result.get("revised_template"):
            result["revised_template"] = normalize_template(result["revised_template"])
        return result

    def import_template(self, filename: str, content: bytes, owner_id: Optional[str] = None) -> str:
        """Convert an uploaded DOCX/PDF/TXT/MD document into a template and store it."""
        if len(content) > 20 * 1024 * 1024:
            raise ValueError("Template files must be under 20 MB")
        validate_upload(filename, content, {".docx", ".pdf", ".txt", ".md"})
        text = _document_text(filename, content)
        if not text.strip():
            raise ValueError("No text could be read from the template file")
        template = normalize_template(self._chat_json(TEMPLATE_IMPORT_SYSTEM, text[:30000], max_tokens=4000))
        return self.store.create_template(template, owner_id)

    # ------------------------------------------------------------------ document generation

    def generate_document(
        self,
        template_id: str,
        drawing_ids: List[str],
        comparison_ids: List[str],
        instructions: str,
        owner_id: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Fill a template from reviewed drawings and comparisons, flagging missing information."""
        template = self.store.get_template(template_id)
        if template is None:
            raise KeyError(template_id)

        sources = []
        for drawing_id in drawing_ids:
            drawing = self.store.get_drawing(drawing_id)
            if drawing and drawing.get("extraction"):
                review = drawing.get("review") or {}
                sources.append({
                    "type": "drawing",
                    "filename": drawing["filename"],
                    "extraction": drawing["extraction"],
                    "review_summary": review.get("summary"),
                    "verdict": review.get("verdict"),
                    "standard": drawing.get("standard"),
                    "findings": review.get("findings"),
                })
        for comparison_id in comparison_ids:
            comparison = self.store.get_comparison(comparison_id)
            if comparison:
                result = dict(comparison["result"])
                result.pop("automated_differences", None)
                sources.append({"type": "comparison", **result})
        if not sources:
            raise ValueError("Select at least one reviewed drawing or comparison as a source")

        sources_text = json.dumps(sources, ensure_ascii=False)[:60000]
        result = self._chat_json(
            GENERATE_SYSTEM,
            f"Template:\n{json.dumps(_template_payload(template))}\n\n"
            f"User instructions: {instructions or 'none'}\n\nSources:\n{sources_text}",
            max_tokens=6000,
            temperature=0,
        )
        title = (result.get("title") or "").strip()
        if not title or title.lower() == template["name"].lower():
            numbers = [
                (s.get("extraction") or {}).get("title_block", {}).get("drawing_number")
                for s in sources if s["type"] == "drawing"
            ]
            numbers = [n for n in numbers if n]
            title = f"{template['name']} – {', '.join(dict.fromkeys(numbers))}" if numbers else template["name"]
        result["title"] = title
        sections, missing = _align_to_template(template, result)
        # Numbers are where a wrong value does real damage; flag any that do not appear in the sources
        unverified = _unverified_numbers(sections, sources_text + " " + (instructions or ""))
        for item in unverified:
            missing.append({
                "section": item["section"],
                "field": item["field"],
                "reason": f"Value '{item['value']}' was not found in the source data – verify against the source",
            })
        document_id = self.store.create_document({
            "template_id": template_id,
            "template_name": template["name"],
            "title": result.get("title") or template["name"],
            "sections": sections,
            "missing_information": missing,
            "source_ids": {"drawings": drawing_ids, "comparisons": comparison_ids},
        }, owner_id)
        return self.store.get_document(document_id)

    def export_markdown(self, document_id: str) -> str:
        document = self.store.get_document(document_id)
        if document is None:
            raise KeyError(document_id)
        lines = [f"# {document['title']}", ""]
        for section in document["sections"]:
            lines += [f"## {section['title']}", ""]
            for name, value in section["fields"].items():
                if "\n|" in f"\n{value}" and value.strip().startswith("|"):
                    lines += [f"**{name}**", "", value, ""]
                else:
                    lines.append(f"**{name}:** {value or '_(missing)_'}")
            if section.get("notes"):
                lines += ["", f"_{section['notes']}_"]
            lines.append("")
        if document.get("missing_information"):
            lines += ["## Missing information", ""]
            lines += [f"- {m['section']} / {m['field']}: {m.get('reason', '')}" for m in document["missing_information"]]
        return "\n".join(lines)

    def export_docx(self, document_id: str) -> bytes:
        import docx

        document = self.store.get_document(document_id)
        if document is None:
            raise KeyError(document_id)
        word = docx.Document()
        word.add_heading(document["title"], level=0)
        for section in document["sections"]:
            word.add_heading(section["title"], level=1)
            for name, value in section["fields"].items():
                rows = _markdown_table_rows(value)
                if rows:
                    word.add_paragraph(name).runs[0].bold = True
                    table = word.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                    table.style = "Table Grid"
                    for r, row in enumerate(rows):
                        for c, cell in enumerate(row):
                            table.cell(r, c).text = cell
                else:
                    paragraph = word.add_paragraph()
                    paragraph.add_run(f"{name}: ").bold = True
                    paragraph.add_run(value or "(missing)")
            if section.get("notes"):
                word.add_paragraph().add_run(section["notes"]).italic = True
        if document.get("missing_information"):
            word.add_heading("Missing information", level=1)
            for item in document["missing_information"]:
                word.add_paragraph(f"{item['section']} / {item['field']}: {item.get('reason', '')}", style="List Bullet")
        buffer = io.BytesIO()
        word.save(buffer)
        return buffer.getvalue()


# ---------------------------------------------------------------------- helpers

def _cad_for_prompt(cad: Optional[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    if not cad:
        return cad
    return {
        "units": cad.get("units"),
        "layers": cad.get("layers"),
        "block_attributes": cad.get("block_attributes"),
        "dimensions": cad.get("dimensions"),
        "texts": [t["text"] for t in cad.get("texts") or []][:200],
    }


def _as_text(value: Any) -> Optional[str]:
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, dict):
        return ", ".join(f"{k}: {v}" for k, v in value.items() if v not in (None, ""))
    if isinstance(value, list):
        return "; ".join(_as_text(item) or "" for item in value)
    return str(value)


def _comparable(extraction: Dict[str, Any]) -> Dict[str, Any]:
    """Extraction without free-text descriptions that vary between runs (they caused false 'changes')."""
    data = {k: v for k, v in extraction.items() if k not in ("unverified_datum_claims", "legibility_issues")}
    data["datum_features"] = sorted(
        str(d.get("letter")) for d in extraction.get("datum_features") or [] if isinstance(d, dict)
    )
    return data


def normalize_template(raw: Dict[str, Any]) -> Dict[str, Any]:
    """Coerce a template (from the UI or the LLM) into the stored schema."""
    sections = []
    for section in raw.get("sections") or []:
        if not isinstance(section, dict) or not str(section.get("title") or "").strip():
            continue
        fields = []
        for field in section.get("fields") or []:
            if isinstance(field, str):
                field = {"name": field}
            name = str(field.get("name") or "").strip()
            if name:
                fields.append({
                    "name": name,
                    "type": field.get("type") if field.get("type") in FIELD_TYPES else "text",
                    "required": bool(field.get("required", True)),
                })
        sections.append({
            "title": str(section["title"]).strip(),
            "guidance": str(section.get("guidance") or "").strip(),
            "fields": fields,
        })
    if not sections:
        raise ValueError("A template needs at least one section")
    return {
        "name": str(raw.get("name") or "Untitled template").strip(),
        "description": str(raw.get("description") or "").strip(),
        "category": str(raw.get("category") or "custom").strip(),
        "sections": sections,
    }


def _template_payload(template: Dict[str, Any]) -> Dict[str, Any]:
    return {key: template[key] for key in ("name", "description", "category", "sections")}


def _align_to_template(template: Dict[str, Any], result: Dict[str, Any]):
    """Keep exactly the template's sections and fields; record required fields left empty."""
    generated = {
        str(s.get("title") or "").strip().lower(): s for s in result.get("sections") or [] if isinstance(s, dict)
    }
    missing = [m for m in result.get("missing_information") or [] if isinstance(m, dict)]
    flagged = {(str(m.get("section")).lower(), str(m.get("field")).lower()) for m in missing}
    sections = []
    for section in template["sections"]:
        source = generated.get(section["title"].lower(), {})
        values = {str(k).strip().lower(): v for k, v in (source.get("fields") or {}).items()}
        fields = {}
        for field in section["fields"]:
            value = values.get(field["name"].lower(), "")
            value = _clean_text(value if isinstance(value, str) else json.dumps(value, ensure_ascii=False))
            fields[field["name"]] = value
            key = (section["title"].lower(), field["name"].lower())
            if field["required"] and not value.strip() and key not in flagged:
                missing.append({"section": section["title"], "field": field["name"],
                                "reason": "Not found in the source documents"})
        sections.append({"title": section["title"], "fields": fields, "notes": source.get("notes") or ""})
    return sections, missing


def _clean_text(value: str) -> str:
    """Remove control characters from model output. DEL (\\x7f) has been observed in place of 'Ø'."""
    value = value.replace("\x7f", "Ø")
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)


# Boundaries are digits/dots only, so numbers right after symbols or letters (Ø9.0, M8, R0.2) are still found
NUMBER_TOKEN = re.compile(r"(?<![\d.])\d+(?:\.\d+)*(?![\d.])")


def _unverified_numbers(sections: List[Dict[str, Any]], sources_text: str) -> List[Dict[str, str]]:
    """Numeric tokens in the generated fields that never occur in the sources (e.g. a corrupted '8.9.0')."""
    source_numbers = set(NUMBER_TOKEN.findall(sources_text))
    flagged = []
    for section in sections:
        for name, value in section["fields"].items():
            for token in set(NUMBER_TOKEN.findall(value or "")):
                if token not in source_numbers and len(token.replace(".", "")) > 1:
                    flagged.append({"section": section["title"], "field": name, "value": token})
    return flagged


def _markdown_table_rows(value: str) -> List[List[str]]:
    lines = [line.strip() for line in (value or "").strip().splitlines() if line.strip()]
    if len(lines) < 2 or not all(line.startswith("|") for line in lines):
        return []
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells if cell):
            continue  # separator row
        rows.append(cells)
    return rows


def _document_text(filename: str, content: bytes) -> str:
    extension = os.path.splitext(filename)[1].lower()
    if extension == ".docx":
        import docx

        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if extension == ".pdf":
        import pymupdf as fitz

        with fitz.open(stream=content, filetype="pdf") as pdf:
            return "\n".join(page.get_text() for page in pdf)
    if extension in (".txt", ".md"):
        return content.decode("utf-8", errors="replace")
    raise ValueError("Template files must be DOCX, PDF, TXT or MD")
