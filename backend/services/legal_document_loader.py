"""Text extraction for legal documents, with OCR for scanned pages and images."""

import base64
import os
from typing import Any, Dict, List

import fitz  # PyMuPDF
from openai import OpenAI

from core.config import settings
from utils.logger import logger
from .usage_tracker import record_usage


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}

# A PDF page with less extractable text than this is treated as scanned and sent to OCR
MIN_TEXT_CHARS_PER_PAGE = 40

OCR_PROMPT = (
    "You are an OCR engine for legal documents. Transcribe ALL text on this page verbatim, "
    "preserving headings, section numbers, clause numbering, tables (as plain text rows), "
    "signature blocks and dates. Do not summarize, translate, or add commentary. "
    "If a word is illegible write [illegible]. If the page has no text, return an empty string."
)


class LegalDocumentLoader:
    """Extracts text from PDF, DOCX, TXT and image files, using vision-model OCR where needed."""

    def __init__(self):
        """Initialize the loader."""
        self.client = OpenAI(api_key=settings.OPENAI_API_KEY, max_retries=5, timeout=180)
        self.model = settings.LEGAL_MODEL

    def load(self, file_path: str, filename: str) -> Dict[str, Any]:
        """
        Extract text from a legal document.

        Args:
            file_path: Path to the file on disk
            filename: Original filename (used to detect the file type)

        Returns:
            Dictionary with the extracted text, page count and which pages were OCR'd

        Raises:
            ValueError: If the file type is unsupported or no text could be extracted
        """
        extension = os.path.splitext(filename)[1].lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type '{extension}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )

        logger.info(f"Loading legal document {filename}")
        if extension == ".pdf":
            result = self._load_pdf(file_path)
        elif extension in IMAGE_EXTENSIONS:
            result = self._load_image(file_path)
        elif extension == ".docx":
            result = self._load_docx(file_path)
        else:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                result = {"text": f.read(), "page_count": 1, "ocr_pages": []}

        if not result["text"].strip():
            raise ValueError("No text could be extracted from the document")

        logger.info(
            f"Extracted {len(result['text'])} characters from {filename} "
            f"({result['page_count']} pages, {len(result['ocr_pages'])} OCR'd)"
        )
        return result

    def _load_pdf(self, file_path: str) -> Dict[str, Any]:
        """Extract text from each PDF page, falling back to OCR for scanned pages."""
        pages: List[str] = []
        ocr_pages: List[int] = []

        with fitz.open(file_path) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text")
                if len(text.strip()) < MIN_TEXT_CHARS_PER_PAGE:
                    if len(ocr_pages) < settings.LEGAL_OCR_MAX_PAGES:
                        text = self._ocr_page(page)
                        ocr_pages.append(page_number)
                    else:
                        logger.warning(f"OCR page limit reached; skipping page {page_number}")
                pages.append(f"[Page {page_number}]\n{text.strip()}")
            page_count = pdf.page_count

        return {"text": "\n\n".join(pages), "page_count": page_count, "ocr_pages": ocr_pages}

    def _load_image(self, file_path: str) -> Dict[str, Any]:
        """OCR a scanned image (multi-page TIFFs are handled page by page)."""
        pages: List[str] = []
        ocr_pages: List[int] = []

        with fitz.open(file_path) as image_doc:
            for page_number, page in enumerate(image_doc, start=1):
                if page_number > settings.LEGAL_OCR_MAX_PAGES:
                    logger.warning(f"OCR page limit reached; skipping page {page_number}")
                    break
                pages.append(f"[Page {page_number}]\n{self._ocr_page(page).strip()}")
                ocr_pages.append(page_number)
            page_count = image_doc.page_count

        return {"text": "\n\n".join(pages), "page_count": page_count, "ocr_pages": ocr_pages}

    def _load_docx(self, file_path: str) -> Dict[str, Any]:
        """Extract paragraphs and table rows from a Word document."""
        import docx

        document = docx.Document(file_path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))

        return {"text": "\n".join(parts), "page_count": 1, "ocr_pages": []}

    def _ocr_page(self, page: "fitz.Page") -> str:
        """Render a page to PNG and transcribe it with the vision-capable LLM."""
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2))
        image_b64 = base64.b64encode(pixmap.tobytes("png")).decode("ascii")

        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": OCR_PROMPT},
                        {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{image_b64}"}},
                    ],
                }
            ],
            temperature=0,
            max_tokens=4000,
        )
        record_usage(self.model, response.usage, "legal-ocr")
        return response.choices[0].message.content or ""
